"""
Build the access layer: do small-scale fishers have recognised rights to fish?

The rest of the atlas says how much small-scale fishers catch and how many of
them there are. It says nothing about whether they have any secure right to keep
fishing - which is the thing Rare's work is about. Two sources answer that, from
opposite ends, and they are kept side by side rather than blended:

  UN SDG indicator 14.b.1     A 1-5 score, every two years, for about 190
                              countries. It is a self-reported policy checklist:
                              does the country have SSF-specific laws, an SSF
                              Guidelines initiative, and an advisory body fishers
                              can sit on? It says nothing about how much catch
                              any of that actually covers.

  Illuminating Hidden Harvests  For 51 countries, how much small-scale catch (in
  governance data               tonnes) falls under each kind of access rule,
  (FAO/Duke/WorldFish)          harvest rule and devolved right. Much richer, but
                              a one-off snapshot (catch 2013-2017, rules
                              2019-2020).

Outputs site/data/access.json, keyed by ISO3.

Sources:
  https://unstats.un.org/SDGAPI/v1/sdg/Series/Data?seriesCode=ER_REG_SSFRAR
      UN SDG Global Database (UN Statistics Division; custodian agency FAO).
      No licence is stated in the API response, so none is claimed here.
  https://unstats.un.org/sdgs/metadata/files/Metadata-14-0b-01.pdf
      The indicator's official metadata (last updated 2024-07-29). The meaning of
      the five levels is quoted from it, because the API itself only says
      "1 lowest to 5 highest".
  https://unstats.un.org/unsd/methodology/m49/overview/
      UN M49 country codes with their ISO3 equivalents - the crosswalk.
  https://www.fao.org/3/cd3915en/IHH_Governance_data.xlsx      (CC BY 4.0)
      Redirects to openknowledge.fao.org; urllib follows redirects by itself.

Every raw download is cached in data/raw/ and reused on the next run. Delete the
cached file to fetch a fresh copy (the SDG series gains a round every two years).

Run:  python3 scripts/fetch_access.py
"""

import json
import os
import sys
import time
import urllib.request
from html.parser import HTMLParser

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import countries as iso_lookup

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(HERE, "data", "raw")
OUT = os.path.join(HERE, "site", "data")

SDG_SERIES = "ER_REG_SSFRAR"
SDG_URL = ("https://unstats.un.org/SDGAPI/v1/sdg/Series/Data"
           f"?seriesCode={SDG_SERIES}&pageSize=5000")
SDG_METADATA_URL = "https://unstats.un.org/sdgs/metadata/files/Metadata-14-0b-01.pdf"
M49_URL = "https://unstats.un.org/unsd/methodology/m49/overview/"
IHH_GOVERNANCE = "https://www.fao.org/3/cd3915en/IHH_Governance_data.xlsx"

SDG_CACHE = os.path.join(RAW, f"sdg_{SDG_SERIES}.json")
SDG_METADATA_CACHE = os.path.join(RAW, "sdg_14b1_metadata.pdf")
M49_CACHE = os.path.join(RAW, "un_m49_overview.html")
IHH_CACHE = os.path.join(RAW, "IHH_Governance_data.xlsx")

# The five levels, quoted from section 2.b of the official metadata PDF above.
# The API's own text only says "1 lowest to 5 highest", which is too thin to put
# on a page. The PDF is a layout document, not data, so the wording is copied
# here once and the PDF is cached beside the other raw files so it can be
# checked. Each level is a band of an underlying 0-1 score that the UN does not
# publish.
SDG_LEVELS = {
    "1": "Very low implementation of instruments for access to resources and markets for small-scale fisheries",
    "2": "Low implementation of instruments for access to resources and markets for small-scale fisheries",
    "3": "Medium implementation of instruments for access to resources and markets for small-scale fisheries",
    "4": "High implementation of instruments for access to resources and markets for small-scale fisheries",
    "5": "Very high implementation of instruments for access to resources and markets for small-scale fisheries",
}
SDG_LEVEL_SCORES = {"1": ">0 to <0.2", "2": "0.2 to <0.4", "3": "0.4 to <0.6",
                    "4": "0.6 to <0.8", "5": "0.8 to 1.0"}
SDG_METADATA_UPDATED = "2024-07-29"

# European Union member states (the 27 since 1 February 2020; source:
# european-union.europa.eu). The API footnotes EU members' scores as built from
# the EU's single response - but it also puts that footnote on some countries
# that are not members (Benin, Gabon, Gambia, Georgia in 2024/2026). Checking
# membership is the only way to tell a real EU-derived score from a stray
# footnote. The United Kingdom was a member until 31 January 2020, so for the
# 2018 round only.
EU_MEMBERS = {"AUT", "BEL", "BGR", "HRV", "CYP", "CZE", "DNK", "EST", "FIN", "FRA",
              "DEU", "GRC", "HUN", "IRL", "ITA", "LVA", "LTU", "LUX", "MLT", "NLD",
              "POL", "PRT", "ROU", "SVK", "SVN", "ESP", "SWE"}
EU_FORMER_LAST_ROUND = {"GBR": 2018}


def _eu_member(iso, year):
    return iso in EU_MEMBERS or year <= EU_FORMER_LAST_ROUND.get(iso, 0)

# What the indicator is built from, also from the metadata PDF (sections 4.a-4.c
# and 6). These are the caveats the page needs to be honest about the score.
SDG_NOTES = [
    "What is scored: three questions from FAO's Code of Conduct for Responsible Fisheries (CCRF) "
    "questionnaire, answered by each country's fisheries administration. (1) laws, regulations, "
    "policies, plans or strategies that specifically target small-scale fisheries, weight 0.4; "
    "(2) ongoing initiatives to implement the SSF Guidelines, weight 0.3; (3) an advisory or "
    "consultative body in which fishers and fish workers can take part, weight 0.3. The 0-1 total "
    "is reported only as a band, 1-5. (Metadata 4.a-4.c.)",
    "It measures whether frameworks exist on paper and are being applied, not how much catch or how "
    "many fishers they cover. It cannot say who actually holds access rights on the water.",
    "Self-reported. The metadata itself warns of self-report bias: countries tend to report 'a better "
    "reality' than the one in place (Metadata section 6). 'Blank' and 'unknown' answers are scored as "
    "'no' (Metadata 4.c).",
    "European Union member states are scored from the EU's single response, not their own; the API "
    "footnote says 'Indicator developped upon responses provided by the EU'. Those rounds are listed "
    "in eu_derived_rounds and counted in latest_eu_derived: a block of identical scores that is really "
    "one answer. The same footnote also appears on some non-members (see eu_footnote_anomalies); those "
    "scores are kept as published and the footnote is treated as a likely error in the source.",
    "Some landlocked countries report on their inland fisheries, although target 14.b is about "
    "marine resources (Metadata 4.b).",
    "A null score means no score was published for that round - never zero. The API's reason is kept "
    "in round_detail[year].missing; missing_reasons counts them.",
    "Regional and other groupings (World, Africa, SIDS, LDCs and so on) are in the same API series; "
    "they are dropped here and listed in dropped_aggregates.",
]


# ---------------------------------------------------------------- downloads

def fetch(url, dest, tries=3):
    """Download once, then reuse. FAO redirects to a bitstream host, so follow."""
    if os.path.exists(dest):
        return dest
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    for attempt in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "ssf-atlas/0.1"})
            with urllib.request.urlopen(req, timeout=120) as r:
                data = r.read()
            break
        except Exception:
            if attempt == tries - 1:
                raise
            time.sleep(2 * (attempt + 1))
    with open(dest, "wb") as f:
        f.write(data)
    return dest


def fetch_sdg():
    """
    The whole 14.b.1 series in one JSON file. It fits in one page today
    (about 1,100 records), but if the UN ever splits it across pages every page
    is fetched and merged before caching, so the cache is always the full series.
    """
    if os.path.exists(SDG_CACHE):
        with open(SDG_CACHE) as f:
            return json.load(f)
    fetch(SDG_URL, SDG_CACHE)
    with open(SDG_CACHE) as f:
        first = json.load(f)
    pages = int(first.get("totalPages") or 1)
    for page in range(2, pages + 1):
        part_path = SDG_CACHE + f".page{page}"
        fetch(f"{SDG_URL}&page={page}", part_path)
        with open(part_path) as f:
            first["data"].extend(json.load(f)["data"])
        os.remove(part_path)
    if pages > 1:
        with open(SDG_CACHE, "w") as f:
            json.dump(first, f)
    return first


# ---------------------------------------------------------------- M49 -> ISO3

class _M49Table(HTMLParser):
    """Reads the English table on the UN M49 overview page into rows of text."""

    def __init__(self):
        super().__init__()
        self.inside = False
        self.rows, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "table" and dict(attrs).get("id") == "downloadTableEN":
            self.inside = True
        elif self.inside and tag == "tr":
            self.row = []
        elif self.inside and tag == "td" and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if not self.inside:
            return
        if tag == "td" and self.cell is not None:
            self.row.append("".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == "table":
            self.inside = False

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)


def m49_to_iso3():
    """
    {m49 number: (ISO3, UN name)} for every country or area the UN lists.

    The SDG API keys countries by M49 number. For countries that is the same as
    the ISO numeric code, but regions (World = 1, Africa = 2 ...) share the same
    number space, so the safe rule is: a code is a country only if the UN's own
    M49 table lists it as one. Everything else is a grouping and gets dropped.
    (fetch_people.py has a helper of the same name, but it returns ISO2 codes from
    the World Bank, not M49 numbers, so it can't be reused here.)
    """
    fetch(M49_URL, M49_CACHE)
    with open(M49_CACHE, encoding="utf-8") as f:
        parser = _M49Table()
        parser.feed(f.read())
    header, body = parser.rows[0], parser.rows[1:]
    col = {name: i for i, name in enumerate(header)}
    i_m49, i_iso, i_name = col["M49 Code"], col["ISO-alpha3 Code"], col["Country or Area"]
    out = {}
    for r in body:
        if len(r) <= max(i_m49, i_iso, i_name):
            continue
        iso, code = r[i_iso].strip().upper(), r[i_m49].strip()
        # Sark has an M49 code but no ISO3; it has no SDG data either.
        if len(iso) == 3 and code.isdigit():
            out[int(code)] = (iso, r[i_name])
    if len(out) < 240:
        raise SystemExit(f"UN M49 table parsed to only {len(out)} countries; page layout changed?")
    return out


# ---------------------------------------------------------------- SDG 14.b.1

def _score(raw):
    """The API sends the band as a string: '1'..'5', or 'NaN'. NaN is a gap, not 0."""
    try:
        v = float(raw)
    except (TypeError, ValueError):
        return None
    if v != v:                     # NaN is the only value not equal to itself
        return None
    if not v.is_integer() or not 1 <= v <= 5:
        raise SystemExit(f"unexpected 14.b.1 value {raw!r}; the series has changed shape")
    return int(v)


def _is_scale_note(text):
    # "Level of implementation: 1 lowest to 5 highest" restates the scale on
    # almost every record; it is kept once at the top, not 560 times.
    return text.strip().lower().startswith("level of implementation")


def _years(ys):
    return ", ".join(str(y) for y in sorted(ys))


def read_sdg(m49):
    data = fetch_sdg()
    records = data["data"]
    if int(data.get("totalElements") or len(records)) != len(records):
        raise SystemExit("14.b.1 download is incomplete; delete the cache and re-run.")

    descriptions = {r["seriesDescription"] for r in records}
    if len(descriptions) != 1:
        raise SystemExit(f"more than one series description: {descriptions}")
    description = descriptions.pop()

    # Code lists for every attribute and dimension, exactly as the API gives them,
    # so a "Q" or "N" on the page can always be spelled out.
    codes = {}
    for block in (data.get("attributes") or []) + (data.get("dimensions") or []):
        codes[block["id"]] = {c["code"]: c["description"] for c in block["codes"]}

    rounds = sorted({int(r["timePeriodStart"]) for r in records})
    latest_round = rounds[-1]

    # Attributes that are the same on every record are reported once.
    uniform = {}
    for key, getter in (("Units", lambda r: r["attributes"].get("Units")),
                        ("Reporting Type", lambda r: r["dimensions"].get("Reporting Type"))):
        vals = {getter(r) for r in records}
        if len(vals) == 1:
            v = vals.pop()
            uniform[key] = {"code": v, "description": codes.get(key, {}).get(v)}

    by_area, dropped, source_tally = {}, {}, {}
    for r in records:
        code = int(r["geoAreaCode"])
        if code not in m49:
            dropped[code] = r["geoAreaName"]
            continue
        by_area.setdefault(code, []).append(r)

    countries = {}
    for code, recs in by_area.items():
        iso, un_name = m49[code]
        if iso in countries:
            raise SystemExit(f"two M49 codes map to {iso}")
        scores, detail, notes_by_text, eu_rounds, eu_odd = {}, {}, {}, [], []
        for r in sorted(recs, key=lambda x: x["timePeriodStart"]):
            year = int(r["timePeriodStart"])
            if str(year) in scores:
                raise SystemExit(f"{iso} has two records for {year}")
            score = _score(r["value"])
            scores[str(year)] = score
            attrs = r.get("attributes") or {}
            notes = [t for t in (r.get("footnotes") or []) if not _is_scale_note(t)]
            rec = {"nature": attrs.get("Nature"),
                   "status": attrs.get("Observation Status")}
            # The record's "source" is one of two long strings; it is tallied once
            # at the top (sources) instead of being repeated on every round.
            src = r.get("source") or "NA"
            source_tally.setdefault(src, {"records": 0, "with_score": 0})
            source_tally[src]["records"] += 1
            source_tally[src]["with_score"] += score is not None
            if score is None:
                # Keep the API's own words for why there is no number.
                rec["missing"] = "; ".join(notes) or codes.get("Observation Status", {}).get(
                    attrs.get("Observation Status"), "not reported")
            if notes and rec.get("missing") != "; ".join(notes):
                rec["footnotes"] = notes
            detail[str(year)] = rec
            for t in notes:
                notes_by_text.setdefault(t, []).append(year)
                if "provided by the EU" in t and score is not None:
                    (eu_rounds if _eu_member(iso, year) else eu_odd).append(year)
        # Every round appears, with null where nothing was published, so the page
        # can draw the gap instead of skipping it.
        for y in rounds:
            scores.setdefault(str(y), None)
        reported = [int(y) for y, s in scores.items() if s is not None]
        latest_year = max(reported) if reported else None
        entry = {
            "name": iso_lookup.display_name(iso, un_name),
            "m49": code,
            "scores": dict(sorted(scores.items())),
            "latest": scores[str(latest_year)] if latest_year else None,
            "latest_year": latest_year,
            "round_detail": dict(sorted(detail.items())),
            "notes": [f"{_years(ys)}: {t}" for t, ys in notes_by_text.items()],
        }
        if eu_rounds:
            entry["eu_derived_rounds"] = sorted(set(eu_rounds))
        if eu_odd:
            entry["eu_footnote_anomaly_rounds"] = sorted(set(eu_odd))
            entry["notes"].append(
                f"{_years(set(eu_odd))}: the API footnotes this score as built from the EU's "
                f"responses, but {entry['name']} is not an EU member state. Most likely a footnote "
                f"error in the source; the score is kept as published.")
        countries[iso] = entry

    # How scores are spread in each round. The page has to say that they bunch
    # at the top, and this is the evidence.
    def spread(year):
        vals = [c["scores"][str(year)] for c in countries.values()]
        d = {str(k): sum(1 for v in vals if v == k) for k in range(1, 6)}
        d["reported"] = sum(1 for v in vals if v is not None)
        d["not_reported"] = sum(1 for v in vals if v is None)
        return d

    by_round = {str(y): spread(y) for y in rounds}
    latest = by_round[str(latest_round)]
    eu_latest = sum(1 for c in countries.values()
                    if latest_round in c.get("eu_derived_rounds", []))
    top = latest["4"] + latest["5"]

    # Why the gaps are gaps, counted over every country-round with no score.
    reasons = {}
    for c in countries.values():
        for rec in c["round_detail"].values():
            if "missing" in rec:
                reasons[rec["missing"]] = reasons.get(rec["missing"], 0) + 1

    return {
        "source": "UN SDG Global Database (UN Statistics Division); custodian agency FAO; "
                  "data from country self-reporting in the CCRF questionnaire",
        "url": SDG_URL,
        "metadata_url": SDG_METADATA_URL,
        "metadata_last_updated": SDG_METADATA_UPDATED,
        "licence": "Not stated in the API response; cite as UN SDG Global Database, indicator 14.b.1.",
        "indicator": "14.b.1",
        "series": SDG_SERIES,
        "description": description,
        "scale_footnote": "Level of implementation: 1 lowest to 5 highest",
        "levels": SDG_LEVELS,
        "level_scores": SDG_LEVEL_SCORES,
        "codes": codes,
        "uniform_attributes": uniform,
        "rounds": rounds,
        "latest_round": latest_round,
        "latest_distribution": latest,
        "latest_share_4_or_5_pct": round(100.0 * top / latest["reported"], 1) if latest["reported"] else None,
        "latest_eu_derived": eu_latest,
        "eu_footnote_anomalies": {iso: c["eu_footnote_anomaly_rounds"] for iso, c in sorted(countries.items())
                                  if c.get("eu_footnote_anomaly_rounds")},
        "sources": source_tally,
        "distribution_by_round": by_round,
        "missing_reasons": dict(sorted(reasons.items(), key=lambda kv: -kv[1])),
        "notes": SDG_NOTES,
        "dropped_aggregates": [{"m49": k, "name": v} for k, v in sorted(dropped.items())],
        "count": len(countries),
        "countries": dict(sorted(countries.items())),
    }


# ---------------------------------------------------------------- IHH governance

# The workbook's own "TYPE" for each field, mapped to the key used in the JSON.
# "Devolved rights (analysis I)" holds two different things - which rights were
# devolved, and how engaged fishers are in management - in two separate blocks
# of the data sheet, so they are kept as two groups here.
GROUP_OF_TYPE = {
    "Access Strategies": "access",
    "Harvesting management measures": "harvest_rules",
    "Devolved rights (analysis I)": "rights",           # *_eng fields go to "engagement"
    "Devolved rights (analysis II)": "rights_breakdown",
}
TOTAL_COL = "National_catch_average (in tonnes)"
RIGHTS_COMBOS = ("omgmt", "oexcl", "otrans", "mgmt_excl", "excl_trans",
                 "mgmt_trans", "mgmt_excl_trans")


def _clean(text):
    return " ".join(str(text).replace("\xa0", " ").split()) if text is not None else None


def _num(v):
    """Tonnes. The workbook writes 'NA' where a tonnage was not available."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return None


def _flag(v):
    if v in (0, 1):
        return bool(v)
    return None


def _pct(part, whole):
    if part is None or not whole:
        return None
    return round(100.0 * part / whole, 1)


def read_ihh_metadata(wb):
    """Field definitions, concept definitions and reference metadata, as written."""
    rows = [tuple(_clean(v) if isinstance(v, str) else v for v in r)
            for r in wb["Metadata"].iter_rows(values_only=True)]
    fields, concepts, reference = {}, {}, {}
    mode = None
    for r in rows:
        r = tuple(r) + (None,) * (4 - len(r))
        a, b, c, d = r[:4]
        if a == "ID" and b == "NAME":
            mode = "fields"
            continue
        if a == "VARIABLE" and b == "DEFINITION":
            mode = "concepts"
            continue
        if a == "Concept" and c == "Identifier":
            mode = "reference"
            continue
        if mode == "fields":
            # The field table ends at the first row that isn't a full ID/NAME/TYPE/
            # DESCRIPTION line.
            if not (a and b and c and d):
                mode = None
                continue
            fields[a] = {"name": b, "type": c, "definition": d}
        elif mode == "concepts":
            if not a or not b:
                mode = None
                continue
            concepts[a] = b
        elif mode == "reference" and c and d is not None:
            reference[c.strip()] = d if not hasattr(d, "isoformat") else d.date().isoformat()
    return fields, concepts, reference


def read_ihh(m49):
    path = fetch(IHH_GOVERNANCE, IHH_CACHE)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # The licence sentence is on the cover sheet; quote it rather than paraphrase.
    licence = None
    for row in wb["Cover"].iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and "Creative Commons" in v:
                licence = _clean(v)
    if not licence or "CC BY 4.0" not in licence:
        raise SystemExit("IHH governance cover sheet no longer states CC BY 4.0; check the licence.")

    fields, concepts, reference = read_ihh_metadata(wb)

    # The metadata lists the total catch as a field too ("national_catch_average",
    # type "Catch"), but in the data sheet it is the column every share is taken
    # of, not one more rule. Take it out of the field list.
    totals = [fid for fid, m in fields.items() if m["type"] == "Catch"]
    if totals != ["national_catch_average"]:
        raise SystemExit(f"expected one 'Catch' field in the metadata, found {totals}")
    total_meta = fields.pop("national_catch_average")
    unknown = {m["type"] for m in fields.values()} - set(GROUP_OF_TYPE)
    if unknown:
        raise SystemExit(f"metadata has field types this script doesn't know: {unknown}")

    # The data sheet's header spans three rows: row 1 names the block (e.g.
    # "SSF Access strategies (they are not mutually exclusive ...)"), row 2 says
    # whether the columns are presence flags or tonnes, row 3 holds the field IDs.
    # The first five columns are merged down all three rows, so their name sits in
    # row 1. Columns are found by name, never by position.
    sheet = wb["IHH_governance_data"]
    rows = list(sheet.iter_rows(values_only=True))
    r1, r3 = rows[0], rows[2]
    header, block_of, block = [], {}, None
    for i in range(len(r1)):
        name = _clean(r3[i]) if i < len(r3) and r3[i] else _clean(r1[i])
        if i >= 5 and r1[i]:
            block = _clean(r1[i])
        header.append(name)
        if name and i >= 5:
            block_of[name] = block
    col = {name: i for i, name in enumerate(header) if name}

    # Each field has a presence column (its ID) and a tonnes column (ID_catch).
    # Management's tonnes column is headed "mgmt_catch.x" - the ".x" looks like a
    # leftover from a join in R, and is noted in the output rather than guessed at.
    catch_col = {}
    for fid in fields:
        if fid not in col:
            raise SystemExit(f"metadata field {fid!r} has no column in the data sheet")
        matches = [h for h in col if h == f"{fid}_catch" or h.startswith(f"{fid}_catch.")]
        if len(matches) != 1:
            raise SystemExit(f"field {fid!r}: expected one tonnes column, found {matches}")
        catch_col[fid] = matches[0]

    def group_of(fid):
        g = GROUP_OF_TYPE[fields[fid]["type"]]
        return "engagement" if g == "rights" and fid.endswith("_eng") else g

    groups = {}
    for fid in fields:
        groups.setdefault(group_of(fid), []).append(fid)

    countries, warnings = {}, []
    facts = {k: [] for k in ("na", "no_info_differs", "eng_without_mgmt", "eng_above_mgmt",
                             "eng_overlap", "blank_name")}
    for row in rows[3:]:
        iso = row[col["Country_ISO3_Code"]]
        if not isinstance(iso, str) or len(iso.strip()) != 3:
            continue
        iso = iso.strip().upper()
        env = _clean(row[col["Inland_marine"]]).lower()
        if env not in ("marine", "inland"):
            raise SystemExit(f"{iso}: unexpected environment {env!r}")
        total = _num(row[col[TOTAL_COL]])
        rec = {"catch_t": total, "flag": _clean(row[col["Flag"]])}
        na_tonnes = []
        for g, ids in groups.items():
            rec[g] = {}
            for fid in ids:
                present = _flag(row[col[fid]])
                tonnes = _num(row[col[catch_col[fid]]])
                if present and tonnes is None:
                    na_tonnes.append(fid)
                rec[g][fid] = {"present": present, "catch_t": tonnes, "pct": _pct(tonnes, total)}

        notes = []
        if na_tonnes:
            notes.append("Present, but tonnage is 'NA' in the workbook for: " + ", ".join(na_tonnes)
                         + ". Their catch_t and pct are null; do not read them as zero.")

        # Analysis II is mutually exclusive, so its parts should add up to the
        # total. Checked here so a broken row is noticed, not silently drawn.
        parts = [rec["rights_breakdown"][f]["catch_t"] for f in groups["rights_breakdown"]]
        if None not in parts and total is not None and abs(sum(parts) - total) > 1:
            notes.append(f"Devolved rights analysis II adds to {sum(parts):,.0f} t, "
                         f"not the {total:,.0f} t total.")

        # Catch under at least one devolved right = the seven right combinations
        # in analysis II. Derived here; null if any of the seven is unknown.
        combos = [rec["rights_breakdown"][f]["catch_t"] for f in RIGHTS_COMBOS]
        if None in combos:
            rec["devolved_any"] = {"catch_t": None, "pct": None}
        else:
            any_t = round(sum(combos), 2)
            rec["devolved_any"] = {"catch_t": any_t, "pct": _pct(any_t, total)}

        # Analysis I's tonnes for each right should equal the analysis II
        # combinations that include it. Where they disagree, say so.
        rb = rec["rights_breakdown"]
        for right, parts_of in (("mgmt", ("omgmt", "mgmt_excl", "mgmt_trans", "mgmt_excl_trans")),
                                ("excl", ("oexcl", "mgmt_excl", "excl_trans", "mgmt_excl_trans")),
                                ("trans", ("otrans", "excl_trans", "mgmt_trans", "mgmt_excl_trans"))):
            a = rec["rights"][right]["catch_t"]
            b = [rb[p]["catch_t"] for p in parts_of]
            if a is not None and None not in b and abs(a - sum(b)) > 1:
                notes.append(f"'{right}' tonnes in analysis I ({a:,.2f} t) differ from the analysis II "
                             f"combinations that include it ({sum(b):,.2f} t), by {sum(b) - a:,.2f} t.")
        if notes:
            rec["notes"] = notes
            warnings.extend(f"{iso} {env}: {n}" for n in notes)

        # Facts the notes below rely on, gathered row by row so the notes are
        # computed from the data rather than written from memory.
        if na_tonnes:
            facts["na"].append(f"{iso} {env}")
        if rec["access"]["no_info"]["catch_t"] != rb["not_gov_info"]["catch_t"]:
            facts["no_info_differs"].append(f"{iso} {env}")
        mg = rec["rights"]["mgmt"]
        for fid, v in rec["engagement"].items():
            if v["present"] and not mg["present"]:
                facts["eng_without_mgmt"].append(f"{iso} {env} {fid}")
            if v["catch_t"] is not None and mg["catch_t"] is not None and v["catch_t"] > mg["catch_t"] + 1:
                facts["eng_above_mgmt"].append(f"{iso} {env} {fid}")
        eng_sum = sum(v["catch_t"] or 0 for v in rec["engagement"].values())
        if mg["catch_t"] and eng_sum > mg["catch_t"] + 1:
            facts["eng_overlap"].append(f"{iso} {env}")

        name = _clean(row[col["Country_name"]])
        entry = countries.setdefault(iso, {"name": None, "marine": None, "inland": None})
        if name:
            entry["name"] = iso_lookup.display_name(iso, name)
        else:
            facts["blank_name"].append(iso)
        if entry[env] is not None:
            raise SystemExit(f"{iso} has two {env} rows")
        entry[env] = rec
    wb.close()

    # A blank name cell (the United Kingdom's, today) falls back to the UN's name.
    un_names = {iso: name for iso, name in m49.values()}
    for iso, e in countries.items():
        if not e["name"]:
            e["name"] = iso_lookup.display_name(iso, un_names.get(iso))

    both = sum(1 for e in countries.values() if e["marine"] and e["inland"])
    marine_only = sum(1 for e in countries.values() if e["marine"] and not e["inland"])
    inland_only = sum(1 for e in countries.values() if e["inland"] and not e["marine"])

    field_out = {}
    for fid, meta in fields.items():
        field_out[fid] = {
            "group": group_of(fid),
            "name": meta["name"],
            "workbook_type": meta["type"],
            "definition": meta["definition"],
            "presence_column": fid,
            "tonnes_column": catch_col[fid],
            "sheet_block": block_of.get(fid),
        }

    group_info = {
        "access": "How access to fish is granted. Not mutually exclusive: one tonne can be "
                  "under a licence and historical use at once, so shares can add to more than 100%.",
        "harvest_rules": "Harvesting management measures in force. Not mutually exclusive.",
        "rights": "Which rights are devolved to fishers (devolved rights, analysis I). Not "
                  "mutually exclusive: a tonne under management and exclusion rights counts in both.",
        "engagement": "How engaged fishers are in management (filed under devolved rights, "
                      "analysis I, in the workbook). Not mutually exclusive.",
        "rights_breakdown": "Devolved rights, analysis II: every tonne is in exactly one category, "
                            "so catch_t adds up to the total and pct to about 100%.",
    }

    notes = [
        f"Coverage: {len(countries)} countries - {both} with marine and inland rows, {marine_only} "
        f"marine only, {inland_only} inland only. The workbook's metadata says: "
        f"\"{reference.get('REF_AREA')}\"",
        "Every 'pct' is computed by this script: that row's tonnes as a share of the same row's "
        "catch_t (marine and inland are never mixed). Tonnes are the workbook's own values.",
        "devolved_any is computed by this script: the sum of the seven analysis II categories in "
        "which at least one right is devolved (everything except no_rights and not_gov_info).",
        f"Time: the workbook says \"{reference.get('COVERAGE_TIME')}\" - so catch is a 2013-2017 "
        f"average and the rules are those recorded in 2019-2020. Flags used: "
        + ", ".join(sorted({e[env]['flag'] for e in countries.values() for env in ('marine', 'inland')
                            if e[env]}))
        + " ('E' = estimate).",
        "Ambiguity: the metadata's description of 'no_rights' is a copy of the high-engagement "
        "one (\"In case there is high engagement ...\"). Its NAME, 'No rights devolved to fishers', "
        "and its place in the mutually exclusive analysis II are taken as its meaning.",
        "Ambiguity: management's tonnes column is headed 'mgmt_catch.x', not 'mgmt_catch'. It is "
        "read as the tonnes for the 'mgmt' field; the '.x' suffix is not explained in the workbook.",
        "Ambiguity: the engagement fields are defined only by their names (e.g. 'High engagement "
        "of fishers in management'); the workbook does not say what they are a share of. "
        + ("In this data they only appear where management rights do and never exceed management "
           "tonnes, so they read as engagement within catch that has management rights. "
           if not facts["eng_without_mgmt"] and not facts["eng_above_mgmt"] else
           "Engagement sometimes appears without management rights: "
           + "; ".join(facts["eng_without_mgmt"] + facts["eng_above_mgmt"]) + ". ")
        + (f"The levels overlap - engagement tonnes add to more than the management tonnes in "
           f"{len(facts['eng_overlap'])} rows ({', '.join(facts['eng_overlap'])}) - so they are "
           f"not a breakdown." if facts["eng_overlap"] else ""),
        "'no_info' (no information on access) and 'not_gov_info' (no information on devolved "
        f"rights) answer different questions; their tonnes differ in {len(facts['no_info_differs'])} "
        f"rows ({', '.join(facts['no_info_differs'])}).",
        "Where a field is present but its tonnage is 'NA' in the workbook, catch_t and pct are null "
        f"and that row's notes list the fields. Rows affected: {', '.join(facts['na']) or 'none'}.",
        "Scope, from the workbook: only formally devolved rights are counted - \"The analysis does "
        "not include governance regimes with informally devolved rights\" - so customary "
        "arrangements are not in these numbers.",
    ]
    if facts["blank_name"]:
        notes.append("The workbook's country name cell is blank for "
                     + ", ".join(facts["blank_name"]) + "; the UN M49 name is used instead.")

    return {
        "source": "Illuminating Hidden Harvests (FAO, Duke University & WorldFish), governance data",
        "url": IHH_GOVERNANCE,
        "citation": "FAO, Duke University & WorldFish. 2023. Illuminating Hidden Harvests - The "
                    "contributions of small-scale fisheries to sustainable development. Rome. "
                    "https://doi.org/10.4060/cc4576en",
        "licence": licence,
        "period": {
            "catch": reference.get("REF_PERIOD"),
            "description": reference.get("COVERAGE_TIME"),
        },
        "metadata_posted": reference.get("META_POSTED"),
        "data_description": reference.get("DATA_DESCR"),
        "unit": reference.get("UNIT_MEASURE"),
        "fields": field_out,
        "concepts": concepts,
        "groups": group_info,
        "output_fields": {
            "catch_t": f"Workbook column '{TOTAL_COL}' ({total_meta['name']}): {total_meta['definition']}",
            "flag": "Workbook column 'Flag'. 'E' = estimate (\"Observation obtained through an estimation methodology\").",
            "<group>.<field>.present": "The field's presence column: true = 1, false = 0.",
            "<group>.<field>.catch_t": "The field's tonnes column; null where the workbook says 'NA'.",
            "<group>.<field>.pct": "catch_t / the row's catch_t x 100, one decimal. Computed here.",
            "devolved_any": "Computed here: catch under at least one devolved right (analysis II).",
            "notes": "Row-level warnings found while reading this row.",
        },
        "notes": notes,
        "warnings": warnings,
        "count": len(countries),
        "countries": dict(sorted(countries.items())),
    }


# ---------------------------------------------------------------- main

def main():
    os.makedirs(OUT, exist_ok=True)

    print("UN M49 codes")
    m49 = m49_to_iso3()
    print(f"  {len(m49)} countries and areas with ISO3")

    print("UN SDG 14.b.1")
    try:
        fetch(SDG_METADATA_URL, SDG_METADATA_CACHE)   # kept for the record; not parsed
    except Exception as e:
        print(f"  ! metadata PDF not cached ({e}); level wording is still the 2024-07-29 text")
    sdg = read_sdg(m49)
    print(f"  {sdg['count']} countries, rounds {sdg['rounds']}; "
          f"{len(sdg['dropped_aggregates'])} groupings dropped")
    d = sdg["latest_distribution"]
    print(f"  {sdg['latest_round']}: " + "  ".join(f"{k}={d[k]}" for k in "12345")
          + f"  reported={d['reported']} not reported={d['not_reported']}"
          + f"  ({sdg['latest_share_4_or_5_pct']}% at 4 or 5; {sdg['latest_eu_derived']} EU-derived)")

    print("IHH governance")
    ihh = read_ihh(m49)
    print(f"  {ihh['count']} countries")
    for w in ihh["warnings"]:
        print(f"  ! {w}")

    with open(os.path.join(OUT, "access.json"), "w") as f:
        json.dump({
            "generated": time.strftime("%Y-%m-%d"),
            "note": ("Two different measures, kept apart. 14.b.1 is a self-reported policy score "
                     "for about 190 countries; IHH is tonnes of catch under each kind of rule for "
                     "51 countries. Neither is a substitute for the other."),
            "sdg14b1": sdg,
            "ihh_governance": ihh,
        }, f, indent=1, ensure_ascii=False)

    # Spot checks a reader can compare with the source by eye.
    print("\n  spot checks")
    for iso in ("PER", "PHL"):
        c = sdg["countries"].get(iso)
        if c:
            print(f"    14.b.1 {iso}: {c['scores']}")
    for iso in ("PHL", "PER"):
        m = (ihh["countries"].get(iso) or {}).get("marine")
        if m:
            rb = m["rights_breakdown"]
            print(f"    IHH {iso} marine: {m['catch_t']:,.0f} t; devolved any "
                  f"{m['devolved_any']['catch_t']:,.0f} t; no rights {rb['no_rights']['catch_t']:,.0f} t; "
                  f"no info {rb['not_gov_info']['catch_t']:,.0f} t ({rb['not_gov_info']['pct']}%)")

    # Which keys will have nothing to attach to on the map.
    with open(os.path.join(OUT, "countries.json")) as f:
        coastal = {c.get("iso3") for c in json.load(f)["countries"] if c.get("iso3")}
    for label, block in (("14.b.1", sdg), ("IHH", ihh)):
        missing = sorted(set(block["countries"]) - coastal)
        print(f"\n  {label}: {len(missing)} ISO3 not in countries.json")
        print("    " + " ".join(missing))


if __name__ == "__main__":
    main()
